import os
import re

from io import BytesIO

from dotenv import load_dotenv

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from pydantic import BaseModel

from openai import OpenAI

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


# =========================================================
# CONFIGURAÇÃO
# =========================================================

load_dotenv()

api_key = os.getenv("OPENAI_API_KEY")

client = OpenAI(api_key=api_key)

app = FastAPI(
    title="Organizador de Esboço de Sermões",
    version="1.0.0"
)


# =========================================================
# CORS
# =========================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# MODELOS DE DADOS
# =========================================================

class PedidoSermao(BaseModel):
    texto_base: str
    titulo: str
    quantidade_topicos: int
    premissas: list[str]


class PedidoSugestao(BaseModel):
    texto_base: str
    titulo: str = ""
    quantidade_topicos: int


class PedidoExportacao(BaseModel):
    esboco: str
    titulo: str = "Esboço de Sermão"
    criador: str = ""
    data_criacao: str = ""


# =========================================================
# PROMPT-MESTRE
# =========================================================

def criar_prompt_sermao(dados: PedidoSermao):

    premissas_formatadas = ""

    for numero, premissa in enumerate(dados.premissas, start=1):
        premissas_formatadas += f"\n{numero}. {premissa}"

    prompt = f"""
Você é um organizador de esboços de sermões bíblicos
de orientação pentecostal e linguagem avivada.

Sua função é ORGANIZAR e DESENVOLVER as premissas
apresentadas pelo pregador.

As premissas fornecidas pelo usuário são os TÓPICOS
PRINCIPAIS do sermão.

Essas premissas podem ter sido escritas diretamente
pelo usuário ou sugeridas anteriormente pela própria IA.
Em ambos os casos, trate-as como os tópicos principais
definitivos deste esboço.


DADOS DO SERMÃO

Texto-base:
{dados.texto_base}

Título:
{dados.titulo}

Quantidade exata de tópicos:
{dados.quantidade_topicos}

Premissas:
{premissas_formatadas}


REGRAS OBRIGATÓRIAS

1. Produza EXATAMENTE {dados.quantidade_topicos} tópicos
principais.

2. Cada premissa recebida deverá corresponder a um tópico
principal.

3. Não elimine nenhuma premissa.

4. Não acrescente tópicos principais além das premissas
recebidas.

5. Preserve a ideia central de cada premissa. Você pode
corrigir discretamente ortografia, concordância e pontuação,
mas não deve substituir a ideia por outra.

6. Para CADA tópico principal, produza EXATAMENTE DUAS
subpremissas.

7. Não produza uma terceira subpremissa.

8. As duas subpremissas devem desenvolver logicamente
o tópico principal.

9. As subpremissas devem ser curtas, objetivas,
memorizáveis e apropriadas para uso no púlpito.

10. Cada subpremissa deverá possuir referência bíblica
correlata que realmente sustente a afirmação apresentada.

11. Não utilize referências apenas por tratarem genericamente
do mesmo assunto.

12. Não invente versículos, personagens, acontecimentos
ou referências bíblicas.

13. Depois das duas subpremissas de cada tópico, escreva
uma pequena Aplicação avivada.

14. A Aplicação avivada deve ser curta, forte, bíblica,
pentecostal e diretamente relacionada ao tópico.

15. A linguagem geral deve ser bíblica, pentecostal,
forte, clara e própria para púlpito.

16. Evite explicações excessivamente longas.

17. O esboço precisa permitir leitura rápida e ser fácil
de memorizar.

18. Ao final, faça uma conclusão curta e um fechamento
avivado relacionado ao texto-base e aos tópicos.

19. Antes do primeiro tópico principal, faça SEMPRE uma
INTRODUÇÃO breve, objetiva e diretamente relacionada ao
tema e ao texto-base.

20. A INTRODUÇÃO deve apresentar resumidamente o contexto
bíblico imediato do texto-base e, quando forem pertinentes
e historicamente seguros, elementos históricos e culturais
que auxiliem na compreensão da passagem.

21. A INTRODUÇÃO deve situar o ouvinte no texto e preparar
naturalmente a passagem para as premissas do sermão.

22. Não transforme a INTRODUÇÃO em estudo acadêmico,
comentário extenso ou exposição excessivamente longa.

23. Não invente informações históricas, culturais,
geográficas, arqueológicas ou contextuais. Quando não
houver segurança suficiente sobre determinada informação,
não a apresente como fato.

24. Não explique estas regras na resposta.

25. NÃO utilize emojis, ícones, desenhos ou símbolos
decorativos em nenhuma parte do esboço.

26. NÃO escreva as palavras "TÍTULO", "TÍTULO DO SERMÃO"
ou expressões semelhantes antes do título.

27. Apresente diretamente o título informado pelo usuário.

28. Utilize negrito Markdown exatamente nos elementos
estruturais indicados no modelo abaixo.

29. Não coloque em negrito todo o texto das subpremissas,
referências ou aplicações. Somente os respectivos rótulos
e os tópicos principais devem receber destaque.

30. Cada subpremissa deve nascer diretamente da lógica da
premissa principal. Não crie pensamentos apenas relacionados
ao tema geral do sermão.

31. Antes de escolher uma referência bíblica, verifique se
o texto citado realmente sustenta a afirmação da subpremissa.

32. Prefira referências bíblicas claras e diretamente
relacionadas à afirmação. Evite associações bíblicas distantes,
forçadas ou apenas temáticas.

33. A referência bíblica não deve servir apenas como
ornamento. Ela deve funcionar como fundamento da subpremissa.

34. Não atribua ao texto-base uma afirmação que ele não faz.

35. Diferencie claramente aquilo que o texto bíblico afirma
daquilo que é aplicação homilética.

36. A Aplicação avivada deve nascer da premissa e das
subpremissas apresentadas. Não introduza na aplicação uma
nova doutrina, acontecimento ou afirmação que não tenha sido
fundamentada no tópico.

37. Evite frases religiosas genéricas que poderiam ser usadas
em qualquer sermão. A aplicação deve ser específica para o
tópico desenvolvido.

38. O modo avivado deve aparecer na força, urgência e
memorização da linguagem, sem sacrificar a precisão bíblica.

39. Quando houver dúvida sobre a adequação de uma referência,
prefira uma referência mais diretamente relacionada ou não
amplie a afirmação.


FORMATO OBRIGATÓRIO DA RESPOSTA

# {dados.titulo}

**Texto-base:** {dados.texto_base}

## INTRODUÇÃO

Introdução breve e objetiva, relacionada ao tema e ao texto-base,
apresentando resumidamente o contexto bíblico e, quando pertinente
e seguro, elementos históricos e culturais que auxiliem na
compreensão da passagem.

1. **{dados.premissas[0] if dados.premissas else "Primeira premissa"}**

**Subpremissa 1:** desenvolvimento curto e lógico.

**Referência bíblica:** referência ou referências adequadas.

**Subpremissa 2:** desenvolvimento curto e lógico.

**Referência bíblica:** referência ou referências adequadas.

**Aplicação avivada:** aplicação curta, forte e memorizável.


Para os demais tópicos, repita EXATAMENTE a mesma estrutura,
utilizando, em ordem, as demais premissas recebidas.

Continue até completar EXATAMENTE
{dados.quantidade_topicos} tópicos.


## CONCLUSÃO

Conclusão curta relacionada ao conjunto do sermão.

**Fechamento avivado:** frase final forte, bíblica,
pentecostal e memorizável.


IMPORTANTE:

O modelo acima demonstra a ESTRUTURA da resposta.

Não escreva explicações sobre o formato.

Não escreva "TÍTULO".

Não utilize emojis.

Não utilize ícones.

Crie apenas UMA introdução, antes do primeiro tópico.

Não crie mais nem menos tópicos do que a quantidade
solicitada.

Cada tópico deve possuir EXATAMENTE duas subpremissas.
"""

    return prompt


# =========================================================
# ROTA PRINCIPAL
# =========================================================

@app.get("/")
def inicio():

    return {
        "status": "online",
        "app": "Organizador de Esboço de Sermões"
    }


# =========================================================
# GERAR ESBOÇO
# =========================================================

@app.post("/gerar-esboco")
def gerar_esboco(dados: PedidoSermao):

    if not api_key or api_key == "SUA_CHAVE_AQUI":

        raise HTTPException(
            status_code=500,
            detail="A chave da API ainda não foi configurada."
        )

    if dados.quantidade_topicos < 1:

        raise HTTPException(
            status_code=400,
            detail="A quantidade de tópicos deve ser maior que zero."
        )

    if len(dados.premissas) != dados.quantidade_topicos:

        raise HTTPException(
            status_code=400,
            detail=(
                "A quantidade de premissas enviadas deve ser "
                "igual à quantidade de tópicos solicitada."
            )
        )

    for premissa in dados.premissas:

        if not premissa.strip():

            raise HTTPException(
                status_code=400,
                detail="Todas as premissas devem ser preenchidas."
            )

    prompt = criar_prompt_sermao(dados)

    try:

        resposta = client.responses.create(
            model="gpt-5.6-luna",
            instructions=(
                "Siga rigorosamente as regras de organização "
                "homilética e de formatação fornecidas. "
                "Faça sempre uma introdução breve antes dos tópicos, "
                "relacionada ao tema e ao texto-base, incluindo contexto "
                "bíblico e, quando pertinente e seguro, elementos "
                "históricos e culturais. "
                "Respeite exatamente a quantidade de tópicos "
                "e produza exatamente duas subpremissas por tópico. "
                "Não utilize emojis ou ícones."
            ),
            input=prompt
        )

        texto = resposta.output_text

        return {
            "sucesso": True,
            "esboco": texto
        }

    except Exception as erro:

        print("ERRO OPENAI:", erro)

        raise HTTPException(
            status_code=500,
            detail="Não foi possível gerar o esboço."
        )


# =========================================================
# SUGERIR PREMISSAS
# =========================================================

@app.post("/sugerir-premissas")
def sugerir_premissas(dados: PedidoSugestao):

    if not api_key or api_key == "SUA_CHAVE_AQUI":

        raise HTTPException(
            status_code=500,
            detail="A chave da API ainda não foi configurada."
        )

    if dados.quantidade_topicos < 1:

        raise HTTPException(
            status_code=400,
            detail="A quantidade de tópicos deve ser maior que zero."
        )

    prompt = f"""
Você é um auxiliar de organização de sermões bíblicos
de orientação pentecostal.

Texto-base:
{dados.texto_base}

Título:
{dados.titulo}

Sugira EXATAMENTE {dados.quantidade_topicos}
premissas para um sermão bíblico pentecostal.

Cada premissa deverá:

- partir diretamente do texto-base;
- ser biblicamente coerente;
- possuir uma ideia própria;
- ser curta;
- ser clara;
- ser forte e apropriada para pregação;
- permitir o desenvolvimento posterior de duas subpremissas;
- contribuir para uma sequência lógica entre os tópicos.

Não desenvolva as premissas.
Não crie subpremissas.
Não apresente referências bíblicas.
Não faça aplicação.
Não faça conclusão.
Não utilize emojis ou ícones.
Não escreva título ou texto-base na resposta.

Retorne SOMENTE as premissas numeradas.
Retorne EXATAMENTE {dados.quantidade_topicos} premissas.
"""

    try:

        resposta = client.responses.create(
            model="gpt-5.6-luna",
            instructions=(
                f"Retorne somente {dados.quantidade_topicos} "
                "premissas numeradas, sem explicações adicionais."
            ),
            input=prompt
        )

        return {
            "sucesso": True,
            "premissas": resposta.output_text
        }

    except Exception as erro:

        print("ERRO OPENAI:", erro)

        raise HTTPException(
            status_code=500,
            detail="Não foi possível sugerir as premissas."
        )


# =========================================================
# EXPORTAR ESBOÇO PARA WORD
# =========================================================

@app.post("/exportar-word")
def exportar_word(dados: PedidoExportacao):

    try:

        documento = Document()

        # -------------------------------------------------
        # MARGENS
        # -------------------------------------------------

        secao = documento.sections[0]

        secao.top_margin = Pt(50)
        secao.bottom_margin = Pt(50)
        secao.left_margin = Pt(60)
        secao.right_margin = Pt(60)

        # -------------------------------------------------
        # FONTE PADRÃO
        # -------------------------------------------------

        estilo = documento.styles["Normal"]

        estilo.font.name = "Calibri"
        estilo.font.size = Pt(12)

        # -------------------------------------------------
        # CONTEÚDO DO ESBOÇO
        # -------------------------------------------------

        linhas = dados.esboco.splitlines()

        for linha in linhas:

            linha = linha.strip()

            if not linha:
                continue

            # Remove Markdown
            linha = linha.replace("**", "")
            linha = linha.replace("###", "")
            linha = linha.replace("##", "")
            linha = linha.replace("#", "")
            linha = linha.strip()

            # ---------------------------------------------
            # TÍTULO
            # ---------------------------------------------

            if linha.upper() == dados.titulo.upper():

                paragrafo = documento.add_paragraph()
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

                run = paragrafo.add_run(linha)

                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(14)

                continue

            # ---------------------------------------------
            # INTRODUÇÃO
            # ---------------------------------------------

            if linha.upper() == "INTRODUÇÃO":

                paragrafo = documento.add_paragraph()

                run = paragrafo.add_run("INTRODUÇÃO")

                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(12)

                continue

            # ---------------------------------------------
            # CONCLUSÃO
            # ---------------------------------------------

            if linha.upper() == "CONCLUSÃO":

                paragrafo = documento.add_paragraph()

                run = paragrafo.add_run("CONCLUSÃO")

                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(12)

                continue

            # ---------------------------------------------
            # TÓPICOS PRINCIPAIS
            # ---------------------------------------------

            if re.match(r"^\d+\.\s+", linha):

                paragrafo = documento.add_paragraph()

                run = paragrafo.add_run(linha)

                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(12)

                continue

            # ---------------------------------------------
            # SUBTÍTULOS EM NEGRITO
            # ---------------------------------------------

            termos_negrito = [
                "Texto-base:",
                "Subpremissa 1:",
                "Subpremissa 2:",
                "Referência bíblica:",
                "Aplicação avivada:",
                "Conclusão:",
                "Fechamento avivado:"
            ]

            encontrou = False

            for termo in termos_negrito:

                if linha.lower().startswith(termo.lower()):

                    paragrafo = documento.add_paragraph()
                    paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

                    inicio = linha[:len(termo)]
                    restante = linha[len(termo):]

                    run_titulo = paragrafo.add_run(inicio)

                    run_titulo.bold = True
                    run_titulo.font.name = "Calibri"
                    run_titulo.font.size = Pt(12)

                    run_texto = paragrafo.add_run(restante)

                    run_texto.font.name = "Calibri"
                    run_texto.font.size = Pt(12)

                    encontrou = True
                    break

            if encontrou:
                continue

            # ---------------------------------------------
            # TEXTO NORMAL
            # ---------------------------------------------

            paragrafo = documento.add_paragraph()
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            run = paragrafo.add_run(linha)

            run.font.name = "Calibri"
            run.font.size = Pt(12)

        # =================================================
        # CRIADOR E DATA DE CRIAÇÃO
        # =================================================

        # Pequeno espaço após o conteúdo
        documento.add_paragraph()

        if dados.criador.strip():

            paragrafo_criador = documento.add_paragraph()
            paragrafo_criador.alignment = WD_ALIGN_PARAGRAPH.LEFT

            rotulo_criador = paragrafo_criador.add_run("Criador: ")
            rotulo_criador.bold = True
            rotulo_criador.font.name = "Calibri"
            rotulo_criador.font.size = Pt(12)

            nome_criador = paragrafo_criador.add_run(
                dados.criador.strip()
            )

            nome_criador.font.name = "Calibri"
            nome_criador.font.size = Pt(12)

        if dados.data_criacao.strip():

            paragrafo_data = documento.add_paragraph()
            paragrafo_data.alignment = WD_ALIGN_PARAGRAPH.LEFT

            rotulo_data = paragrafo_data.add_run(
                "Data de criação: "
            )

            rotulo_data.bold = True
            rotulo_data.font.name = "Calibri"
            rotulo_data.font.size = Pt(12)

            texto_data = paragrafo_data.add_run(
                dados.data_criacao.strip()
            )

            texto_data.font.name = "Calibri"
            texto_data.font.size = Pt(12)

        # =================================================
        # CRIAR ARQUIVO EM MEMÓRIA
        # =================================================

        arquivo = BytesIO()

        documento.save(arquivo)

        arquivo.seek(0)

        # =================================================
        # NOME DO ARQUIVO
        # =================================================

        nome_arquivo = re.sub(
            r'[\\/*?:"<>|]',
            "",
            dados.titulo
        ).strip()

        if not nome_arquivo:
            nome_arquivo = "Esboco_de_Sermao"

        nome_arquivo = nome_arquivo.replace(" ", "_") + ".docx"

        # =================================================
        # ENVIAR WORD PARA O NAVEGADOR
        # =================================================

        return StreamingResponse(
            arquivo,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            headers={
                "Content-Disposition":
                    f'attachment; filename="{nome_arquivo}"'
            }
        )

    except Exception as erro:

        print("ERRO AO EXPORTAR WORD:", erro)

        raise HTTPException(
            status_code=500,
            detail="Não foi possível gerar o arquivo Word."
        )